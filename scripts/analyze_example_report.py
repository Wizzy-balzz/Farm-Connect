import docx

doc_path = r"d:\MWTEL\myi-react-app\example-report\Store_Management_System_Report_Revised.docx"
doc = docx.Document(doc_path)

print(f"Total Paragraphs: {len(doc.paragraphs)}")
print(f"Total Tables: {len(doc.tables)}")

paragraphs_info = []
for i, p in enumerate(doc.paragraphs):
    if p.text.strip():
        style_name = p.style.name if p.style else "Normal"
        runs_info = []
        for r in p.runs:
            runs_info.append({
                "text": r.text,
                "bold": r.bold,
                "italic": r.italic,
                "font_name": r.font.name,
                "font_size": r.font.size.pt if r.font.size else None,
                "color": str(r.font.color.rgb) if (r.font and r.font.color and r.font.color.rgb) else None
            })
        paragraphs_info.append({
            "idx": i,
            "style": style_name,
            "alignment": str(p.alignment),
            "text": p.text[:120],
            "full_text": p.text,
            "runs": runs_info
        })

print("\n--- ALL PARAGRAPHS OVERVIEW ---")
with open("scripts/example_doc_structure.txt", "w", encoding="utf-8") as f:
    f.write(f"Total Paragraphs: {len(doc.paragraphs)}\n")
    f.write(f"Total Tables: {len(doc.tables)}\n\n")
    for p in paragraphs_info:
        f.write(f"[{p['idx']}] Style: {p['style']} | Align: {p['alignment']}\nText: {p['full_text']}\n")
        f.write("-" * 50 + "\n")

print("Saved full paragraph breakdown to scripts/example_doc_structure.txt")

print("\n--- FIRST 40 PARAGRAPHS ---")
for p in paragraphs_info[:40]:
    print(f"[{p['idx']}] Style: {p['style']} | Align: {p['alignment']} | Text: {p['full_text']}")

print("\n--- TABLES SUMMARY ---")
for idx, table in enumerate(doc.tables):
    print(f"Table {idx}: Rows={len(table.rows)}, Cols={len(table.columns)}")
    if len(table.rows) > 0:
        first_row = [cell.text.strip().replace('\n', ' ') for cell in table.rows[0].cells]
        print(f"   Header: {first_row[:5]}")
