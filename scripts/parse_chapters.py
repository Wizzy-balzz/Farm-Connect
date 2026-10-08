import docx
import os

doc_path = r"d:\MWTEL\myi-react-app\example-report\Store_Management_System_Report_Revised.docx"
doc = docx.Document(doc_path)

print("--- COVER & FRONT MATTER ---")
for p in doc.paragraphs[:35]:
    if p.text.strip():
        print(f"[{p.alignment}] {p.text}")

print("\n--- CHAPTER HEADINGS AND SUBHEADINGS ---")
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if txt.startswith("CHAPTER") or txt.startswith("TABLE OF CONTENTS") or txt.startswith("BONAFIDE") or (p.style and 'Heading' in p.style.name):
        print(f"Line {i} | Style: {p.style.name if p.style else 'None'} | Text: {txt}")

# Check images in document
images_count = 0
for rel in doc.part.rels.values():
    if "image" in rel.target_ref:
        images_count += 1
print(f"\nTotal embedded images in docx: {images_count}")

# Check headers and footers
for s_idx, section in enumerate(doc.sections):
    print(f"\nSection {s_idx} Header:")
    for hp in section.header.paragraphs:
        if hp.text.strip(): print("  HDR:", hp.text)
    print(f"Section {s_idx} Footer:")
    for fp in section.footer.paragraphs:
        if fp.text.strip(): print("  FTR:", fp.text)
