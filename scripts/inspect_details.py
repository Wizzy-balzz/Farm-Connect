import docx
import json

doc_path = r"d:\MWTEL\myi-react-app\example-report\Store_Management_System_Report_Revised.docx"
doc = docx.Document(doc_path)

print("=== DOCUMENT SETTINGS & STYLES ===")
print("Sections count:", len(doc.sections))
for i, sec in enumerate(doc.sections):
    print(f"Sec {i}: Top={sec.top_margin.pt}pt, Bottom={sec.bottom_margin.pt}pt, Left={sec.left_margin.pt}pt, Right={sec.right_margin.pt}pt, PageW={sec.page_width.pt}pt, PageH={sec.page_height.pt}pt")
    print(f"       Different First Page: {sec.different_first_page_header_footer}")

full_text_dump = []
for p in doc.paragraphs:
    full_text_dump.append({
        "style": p.style.name if p.style else None,
        "text": p.text,
        "align": str(p.alignment)
    })

with open("scripts/full_example_dump.json", "w", encoding="utf-8") as f:
    json.dump(full_text_dump, f, indent=2)

print("Saved full example dump to scripts/full_example_dump.json")

# Let's inspect sample tables
print("\n=== SAMPLE TABLES DETAILED INSPECTION ===")
for i in range(min(5, len(doc.tables))):
    t = doc.tables[i]
    print(f"\n--- TABLE {i} ---")
    for r_idx, row in enumerate(t.rows):
        cells_txt = [c.text.strip().replace('\n', ' ') for c in row.cells]
        print(f"Row {r_idx}: {cells_txt}")
