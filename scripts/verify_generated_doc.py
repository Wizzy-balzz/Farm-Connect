import docx
import zipfile

doc_path = r"d:\MWTEL\myi-react-app\example-report\FarmConnect_Project_Report_Revised.docx"
doc = docx.Document(doc_path)

print(f"Generated Document Paragraphs: {len(doc.paragraphs)}")
print(f"Generated Document Tables: {len(doc.tables)}")
print(f"Generated Document Sections: {len(doc.sections)}")

text_all = "\n".join([p.text for p in doc.paragraphs])
print("Contains student name (BALASUBRAMANIYAM A):", "BALASUBRAMANIYAM A" in text_all)
print("Contains register number (24104023):", "24104023" in text_all)

with zipfile.ZipFile(doc_path, 'r') as z:
    media_files = [f for f in z.namelist() if 'media' in f]
    print(f"Embedded media files count in generated docx: {len(media_files)}")

# Check chapter headings present
chapters_found = [p.text for p in doc.paragraphs if "CHAPTER" in p.text]
print("Chapters found:", chapters_found)
