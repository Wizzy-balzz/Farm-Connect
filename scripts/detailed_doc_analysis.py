import docx
import json

doc_path = r"d:\MWTEL\myi-react-app\example-report\Store_Management_System_Report_Revised.docx"
doc = docx.Document(doc_path)

out_lines = []

out_lines.append(f"TOTAL PARAGRAPHS: {len(doc.paragraphs)}")
out_lines.append(f"TOTAL TABLES: {len(doc.tables)}")
out_lines.append(f"TOTAL SECTIONS: {len(doc.sections)}")

for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if not txt:
        continue
    
    style_name = p.style.name if p.style else "Normal"
    align = str(p.alignment)
    
    # Extract font details of first run if present
    font_name = None
    font_size = None
    bold = None
    if p.runs:
        font_name = p.runs[0].font.name
        font_size = p.runs[0].font.size.pt if p.runs[0].font.size else None
        bold = p.runs[0].bold

    out_lines.append(f"P[{i}] | Style={style_name} | Align={align} | Font={font_name}, Size={font_size}, Bold={bold}")
    out_lines.append(f"     Text: {txt[:150]}...")

with open("scripts/paragraph_analysis.txt", "w", encoding="utf-8") as f:
    f.write("\n".join(out_lines))

print("Paragraph analysis written to scripts/paragraph_analysis.txt")

# Let's also inspect all table structures
table_info = []
for t_idx, t in enumerate(doc.tables):
    rows_data = []
    for r in t.rows:
        row_txts = [c.text.strip().replace('\n', ' ') for c in r.cells]
        rows_data.append(row_txts)
    table_info.append({
        "table_index": t_idx,
        "rows_count": len(t.rows),
        "cols_count": len(t.columns) if t.rows else 0,
        "header": rows_data[0] if rows_data else [],
        "sample_rows": rows_data[1:3] if len(rows_data) > 1 else []
    })

with open("scripts/table_analysis.json", "w", encoding="utf-8") as f:
    json.dump(table_info, f, indent=2)

print("Table analysis written to scripts/table_analysis.json")
