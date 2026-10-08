import docx
from docx.oxml.ns import nsdecls
from docx.oxml import parse_xml

doc_path = r"d:\MWTEL\myi-react-app\example-report\Store_Management_System_Report_Revised.docx"
doc = docx.Document(doc_path)

print("=== PARAGRAPH STYLES ===")
styles_used = set()
for p in doc.paragraphs:
    if p.style:
        styles_used.add(p.style.name)
print("Styles used in document:", styles_used)

print("\n=== SAMPLE TABLE XML BORDERS & SHADING ===")
if len(doc.tables) > 0:
    t0 = doc.tables[0]
    cell0 = t0.rows[0].cells[0]
    print("Cell XML:", cell0._tc.xml[:300])

print("\n=== SECTION MARGINS & PAGE SIZE ===")
for sec in doc.sections:
    print(f"Top: {sec.top_margin.inches}in, Bottom: {sec.bottom_margin.inches}in, Left: {sec.left_margin.inches}in, Right: {sec.right_margin.inches}in")
    print(f"Page Width: {sec.page_width.inches}in, Page Height: {sec.page_height.inches}in")
